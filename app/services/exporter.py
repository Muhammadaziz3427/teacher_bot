"""Excel (XLSX) export of everything the teacher may need for records."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from ..config import settings
from ..i18n import category_label
from ..utils import fmt_date_short, now
from . import attendance as attendance_service
from . import reports as reports_service
from . import scoring
from . import submissions as submission_service

HEADER_FILL = PatternFill("solid", start_color="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(sheet, row: int = 1) -> None:
    for cell in sheet[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(sheet) -> None:
    for column in sheet.columns:
        longest = max((len(str(cell.value or "")) for cell in column), default=8)
        letter = get_column_letter(column[0].column)
        sheet.column_dimensions[letter].width = min(max(longest + 2, 10), 48)


async def build_workbook(
    session: AsyncSession, chat_id: int, period: str = "weekly", lang: str = "bi"
) -> Path:
    """Create an XLSX with summary, attendance, homework and mistake sheets."""
    start, end = reports_service.period_bounds(period)
    stats_list = await reports_service.collect_all(session, chat_id, period)
    start_dt, end_dt = reports_service.as_datetimes(start, end)

    workbook = Workbook()

    # ---- summary -------------------------------------------------------
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Student", "Attendance %", "Lessons", "Submitted", "Assigned",
                    "On time", "Late", "Missing", "Average", "Points", "Level"])
    for stats in stats_list:
        summary.append([
            stats.student.full_name,
            stats.attendance_pct,
            stats.attendance.get("total", 0),
            stats.submitted,
            stats.assigned,
            stats.ontime,
            stats.late,
            stats.missing,
            stats.avg,
            stats.points,
            stats.level,
        ])
    _style_header(summary)
    _autosize(summary)

    # ---- attendance ----------------------------------------------------
    attendance_sheet = workbook.create_sheet("Attendance")
    attendance_sheet.append(["Student", "Date", "Status", "Marked at"])
    for stats in stats_list:
        rows = await attendance_service.for_student(
            session, chat_id, stats.student.tg_id, start, end
        )
        for row in rows:
            attendance_sheet.append([
                stats.student.full_name,
                fmt_date_short(row.day),
                row.status,
                row.marked_at.strftime("%d.%m.%Y %H:%M") if row.marked_at else "",
            ])
    _style_header(attendance_sheet)
    _autosize(attendance_sheet)

    # ---- homework ------------------------------------------------------
    homework_sheet = workbook.create_sheet("Homework")
    homework_sheet.append(["Student", "Homework #", "Title", "Skill", "Due",
                           "Submitted at", "Late", "Score", "Grade", "Points"])
    for stats in stats_list:
        rows = await submission_service.for_student(
            session, chat_id, stats.student.tg_id, start_dt, end_dt
        )
        for row in rows:
            homework_sheet.append([
                stats.student.full_name,
                row.homework_id,
                "",
                "",
                "",
                row.submitted_at.strftime("%d.%m.%Y %H:%M") if row.submitted_at else "",
                "yes" if row.is_late else "no",
                row.ai_score,
                row.ai_grade,
                row.points_awarded,
            ])
    _style_header(homework_sheet)
    _autosize(homework_sheet)

    # ---- mistakes ------------------------------------------------------
    mistakes = workbook.create_sheet("Mistakes")
    mistakes.append(["Category", "Count (whole class)"])
    all_subs = await submission_service.between(session, chat_id, start_dt, end_dt)
    for category, count in submission_service.mistake_categories(all_subs).items():
        mistakes.append([category_label(category, lang), count])
    _style_header(mistakes)
    _autosize(mistakes)

    # ---- points ledger -------------------------------------------------
    points = workbook.create_sheet("Points")
    points.append(["Student", "Date", "Kind", "Points", "Reason"])
    for stats in stats_list:
        for event in await scoring.recent(session, chat_id, stats.student.tg_id, limit=500):
            if start_dt <= event.created_at <= end_dt:
                points.append([
                    stats.student.full_name,
                    event.created_at.strftime("%d.%m.%Y %H:%M"),
                    event.kind,
                    event.points,
                    event.reason,
                ])
    _style_header(points)
    _autosize(points)

    target_dir = settings.exports_path
    target_dir.mkdir(parents=True, exist_ok=True)
    stamp = now().strftime("%Y%m%d_%H%M%S")
    path = target_dir / f"report_{period}_{stamp}.xlsx"
    workbook.save(path)
    return path
